"""Real-browser integration test for `extract_slide_elements()`.

Unlike the other editable-PPTX tests (which mock Playwright/the browser
entirely), this one drives an actual installed Chromium-family browser
via Playwright, end to end, against a real Blattwerk-rendered presentation
document -- the only test in the suite that proves the DOM-extraction
mechanism itself (not just the surrounding orchestration/conversion
logic) actually works.

Skips cleanly (not a failure) when either `playwright` isn't installed or
no supported browser is found -- both expected, since Playwright is an
optional dependency (`requirements-editable-pptx.txt`), not something
every dev/CI environment is guaranteed to have. Manually verified once
against real Chrome during implementation (`find_chromium_executable()`
resolved to an installed Chrome, Playwright launched it via
`executable_path`, correctly extracted text -- including text split
around an inline MathJax formula -- and screenshotted a QR-style image
block; `:::table` becomes a real cell-based PPTX table since B4, not a
screenshot).
"""

from pathlib import Path

import pytest

from app.core.blatt_kern_io_build import build_worksheet
from app.core.blatt_kern_pptx_export import _slide_size_emu
from app.core.build_requests import WorksheetDesignOptions

try:
    import playwright  # noqa: F401

    _PLAYWRIGHT_INSTALLED = True
except ImportError:
    _PLAYWRIGHT_INSTALLED = False


def _browser_available() -> bool:
    from app.core.blatt_kern_io_pdf import find_chromium_executable

    return bool(find_chromium_executable())


pytestmark = pytest.mark.skipif(
    not _PLAYWRIGHT_INSTALLED,
    reason="playwright nicht installiert (siehe requirements-editable-pptx.txt) -- optionale Abhaengigkeit",
)

_SAMPLE_MARKDOWN = """---
Titel: Testfolien
Fach: Mathematik
Thema: PPTX-Extraktion
mode: presentation
---

# Erste Folie

:::task title="Kernaussage"
Formuliere eine klare Aussage in **einem** Satz zur Funktion $f(x) = x^2$.
:::

:::info type=note
Ein kurzer Hinweistext zum Nachdenken.
:::

-+

:::table
cells:
  - ["A", "B"]
  - ["1", "2"]
:::
"""


@pytest.fixture
def rendered_presentation_html(tmp_path) -> Path:
    if not _browser_available():
        pytest.skip("kein installierter Chromium-Browser gefunden (find_chromium_executable())")

    md_path = tmp_path / "praesentation.md"
    md_path.write_text(_SAMPLE_MARKDOWN, encoding="utf-8")
    html_path = tmp_path / "praesentation.html"
    build_worksheet(
        str(md_path), str(html_path), page_format="presentation_16_9",
        **WorksheetDesignOptions("indigo", "segoe", "normal").as_kwargs(),
    )
    return html_path


def test_extract_slide_elements_finds_two_slides(rendered_presentation_html):
    from app.core.blatt_kern_pptx_export_editable import extract_slide_elements

    width_emu, height_emu = _slide_size_emu("presentation_16_9")
    results = extract_slide_elements(rendered_presentation_html, width_emu, height_emu, mathjax_wait_ms=3000)

    assert len(results) == 2
    assert all(r.elements is not None for r in results)  # Stage A, no fallback needed


def test_extract_slide_elements_captures_full_sentence_around_inline_formula(rendered_presentation_html):
    # Regression for the mixed-content gap found during implementation:
    # plain text nodes sitting next to an <mjx-container> formula inside
    # the same paragraph must not silently vanish.
    from app.core.blatt_kern_pptx_export_editable import extract_slide_elements

    width_emu, height_emu = _slide_size_emu("presentation_16_9")
    results = extract_slide_elements(rendered_presentation_html, width_emu, height_emu, mathjax_wait_ms=3000)

    all_text = " ".join(
        run.text
        for result in results
        for element in (result.elements or [])
        if element.kind == "text"
        for run in (element.runs or [])
    )
    assert "Formuliere eine klare Aussage in" in all_text
    assert "einem" in all_text
    assert "Satz zur Funktion" in all_text


_MIXED_FORMATTING_MARKDOWN = """---
Titel: Testfolien
Fach: Mathematik
Thema: PPTX-Inline-Formatierung
mode: presentation
---

:::task title="Formatierter Text"
Zuerst **fett gedruckt** dann normaler Text dann *kursiv gesetzt* am Ende.
:::
"""


@pytest.fixture
def rendered_mixed_formatting_html(tmp_path) -> Path:
    if not _browser_available():
        pytest.skip("kein installierter Chromium-Browser gefunden (find_chromium_executable())")

    md_path = tmp_path / "mixed_formatting.md"
    md_path.write_text(_MIXED_FORMATTING_MARKDOWN, encoding="utf-8")
    html_path = tmp_path / "mixed_formatting.html"
    build_worksheet(
        str(md_path), str(html_path), page_format="presentation_16_9",
        **WorksheetDesignOptions("indigo", "segoe", "normal").as_kwargs(),
    )
    return html_path


def test_extract_slide_elements_captures_mixed_bold_italic_text_as_ordered_runs(rendered_mixed_formatting_html):
    # B3: "**fett gedruckt** dann normaler Text dann *kursiv gesetzt* am
    # Ende" must come back as ONE text element with several `TextRun`s
    # (not three unrelated text elements, and not one run carrying the
    # whole paragraph's flat style) -- in document order, each run's own
    # bold/italic flag correct, with no lost/duplicated text or whitespace
    # across the run boundaries.
    from app.core.blatt_kern_pptx_export_editable import extract_slide_elements

    width_emu, height_emu = _slide_size_emu("presentation_16_9")
    results = extract_slide_elements(rendered_mixed_formatting_html, width_emu, height_emu, mathjax_wait_ms=1000)

    assert len(results) == 1
    text_elements = [el for el in (results[0].elements or []) if el.kind == "text"]
    matching = [el for el in text_elements if el.runs and "fett gedruckt" in "".join(r.text for r in el.runs)]
    assert len(matching) == 1
    runs = matching[0].runs
    assert len(runs) >= 3

    full_text = "".join(run.text for run in runs)
    assert full_text == "Zuerst fett gedruckt dann normaler Text dann kursiv gesetzt am Ende."

    bold_runs = [run for run in runs if run.bold]
    italic_runs = [run for run in runs if run.italic]
    assert bold_runs and all(run.text.strip() == "fett gedruckt" for run in bold_runs)
    assert italic_runs and all(run.text.strip() == "kursiv gesetzt" for run in italic_runs)
    # The bold run must precede the italic run (document order preserved).
    assert runs.index(bold_runs[0]) < runs.index(italic_runs[0])


def test_extract_slide_elements_table_block_becomes_a_real_table(rendered_presentation_html):
    # B4: `:::table` used to become a single screenshot image -- it is now
    # a real `kind="table"` element with actual cell content/geometry
    # (`cells: [["A", "B"], ["1", "2"]]` from `_SAMPLE_MARKDOWN`), not a
    # picture.
    from app.core.blatt_kern_pptx_export_editable import extract_slide_elements

    width_emu, height_emu = _slide_size_emu("presentation_16_9")
    results = extract_slide_elements(rendered_presentation_html, width_emu, height_emu, mathjax_wait_ms=3000)

    table_elements = [el for el in (results[1].elements or []) if el.kind == "table"]
    assert len(table_elements) == 1
    table = table_elements[0].table
    assert table is not None
    # `:::table` (no explicit `rows=`) defaults to 4 rows (`_render_table_answer`'s
    # own default) even though the `cells:` payload only fills the first
    # two -- the remaining rows come back as real, empty-text cells, not
    # be dropped.
    assert table.cols == 2 and table.rows >= 2
    cells_by_text = {cell.text: cell for cell in table.cells if cell.text}
    assert set(cells_by_text) == {"A", "B", "1", "2"}
    assert cells_by_text["A"].row == 0 and cells_by_text["A"].col == 0
    assert cells_by_text["B"].row == 0 and cells_by_text["B"].col == 1
    assert cells_by_text["1"].row == 1 and cells_by_text["1"].col == 0
    assert cells_by_text["2"].row == 1 and cells_by_text["2"].col == 1


_MERGED_TABLE_MARKDOWN = """---
Titel: Testfolien
Fach: Mathematik
Thema: PPTX-Tabelle
mode: presentation
---

:::table
cells:
  - [{text: "Ueberschrift", colspan: 2}, "C"]
  - ["1", "2", "3"]
:::
"""


def test_extract_slide_elements_table_with_merged_cell_roundtrips_through_real_pptx(tmp_path):
    # Full pipeline (markdown -> build_presentation_pptx(editable=True) ->
    # real .pptx read back via python-pptx), not just the extraction step
    # -- proves a `:::table` with an actual merged (colspan) cell survives
    # end to end: correct cell texts AND the merge itself, matching the
    # plan's explicit B4 verification requirement.
    if not _browser_available():
        pytest.skip("kein installierter Chromium-Browser gefunden (find_chromium_executable())")

    from pptx import Presentation as PptxPresentation

    from app.core.blatt_kern_pptx_export import build_presentation_pptx

    md_path = tmp_path / "merged_table.md"
    md_path.write_text(_MERGED_TABLE_MARKDOWN, encoding="utf-8")
    out_path = tmp_path / "merged_table.pptx"
    diagnostics = []

    build_presentation_pptx(
        md_path, out_path, page_format="presentation_16_9",
        design=WorksheetDesignOptions("indigo", "segoe", "normal"),
        editable=True, diagnostics_out=diagnostics,
    )

    assert out_path.exists()
    assert not any(d.code in {"PPTX001", "PPTX002"} for d in diagnostics)

    prs = PptxPresentation(str(out_path))
    table_shapes = [shape for shape in prs.slides[0].shapes if shape.has_table]
    assert len(table_shapes) == 1
    table = table_shapes[0].table

    assert table.cell(0, 0).text == "Ueberschrift"
    assert table.cell(0, 0).is_spanned is False
    assert table.cell(0, 0).span_width == 2
    assert table.cell(0, 1).is_spanned is True
    assert table.cell(0, 2).text == "C"
    assert [table.cell(1, col).text for col in range(3)] == ["1", "2", "3"]


_CHROME_ONLY_MARKDOWN = """---
Titel: Testfolien
Fach: Mathematik
Thema: PPTX-Chrome
mode: presentation
---

# Erste Folie

:::task title="Nur Text"
Ein einfacher, rein textueller Aufgabentext ohne Bilder/Tabellen/Formeln.
:::
"""


@pytest.fixture
def rendered_chrome_only_html(tmp_path) -> Path:
    if not _browser_available():
        pytest.skip("kein installierter Chromium-Browser gefunden (find_chromium_executable())")

    md_path = tmp_path / "chrome_only.md"
    md_path.write_text(_CHROME_ONLY_MARKDOWN, encoding="utf-8")
    html_path = tmp_path / "chrome_only.html"
    build_worksheet(
        str(md_path), str(html_path), page_format="presentation_16_9",
        **WorksheetDesignOptions("indigo", "segoe", "normal").as_kwargs(),
    )
    return html_path


def test_extract_slide_elements_captures_slide_chrome_as_two_images(rendered_chrome_only_html):
    # Confirms `data-block-type="chrome"` (blatt_kern_layout_presentation.py)
    # is actually picked up by the extractor, not silently invisible.
    # Does NOT assume chrome is the ONLY source of images on the slide --
    # a plain `:::task` already renders its own work-symbol icon as a
    # separate image (discovered while writing this test), so this looks
    # specifically for wide-and-short "chrome-shaped" images (near full
    # slide width, well under a quarter of slide height) rather than
    # asserting an exact total image count.
    from app.core.blatt_kern_pptx_export_editable import extract_slide_elements

    width_emu, height_emu = _slide_size_emu("presentation_16_9")
    results = extract_slide_elements(rendered_chrome_only_html, width_emu, height_emu, mathjax_wait_ms=1000)

    assert len(results) == 1
    images = [el for el in (results[0].elements or []) if el.kind == "image"]
    chrome_shaped = [
        img for img in images
        if img.image_bytes and img.width_emu > width_emu * 0.8 and img.height_emu < height_emu * 0.25
    ]
    assert len(chrome_shaped) >= 2

    top_region = min(chrome_shaped, key=lambda img: img.top_emu)
    bottom_region = max(chrome_shaped, key=lambda img: img.top_emu)
    assert top_region.top_emu < height_emu * 0.3
    assert bottom_region.top_emu > height_emu * 0.6


_ONE_PX_PNG = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
    b"\x08\x02\x00\x00\x00\x90wS\xde\x00\x00\x00\x0cIDATx\x9cc\xf8\xcf\xc0"
    b"\x00\x00\x03\x01\x01\x00\x18\xdd\x8d\xb0\x00\x00\x00\x00IEND\xaeB`\x82"
)

_IMAGE_MARKDOWN = """---
Titel: Testfolien
Fach: Mathematik
Thema: PPTX-Bild-Originalbytes
mode: presentation
---

:::task title="Mit Bild"
Hier ein eingebettetes Bild: ![Test](original.png)
:::
"""


@pytest.fixture
def rendered_image_html(tmp_path) -> Path:
    if not _browser_available():
        pytest.skip("kein installierter Chromium-Browser gefunden (find_chromium_executable())")

    (tmp_path / "original.png").write_bytes(_ONE_PX_PNG)
    md_path = tmp_path / "mit_bild.md"
    md_path.write_text(_IMAGE_MARKDOWN, encoding="utf-8")
    html_path = tmp_path / "mit_bild.html"
    build_worksheet(
        str(md_path), str(html_path), page_format="presentation_16_9",
        **WorksheetDesignOptions("indigo", "segoe", "normal").as_kwargs(),
    )
    return html_path


def test_extract_slide_elements_uses_original_image_bytes_not_a_screenshot(rendered_image_html):
    # B2: a real <img> reference must be extracted using its ORIGINAL
    # source bytes (byte-identical to the file on disk), not a re-encoded
    # browser screenshot of the rendered pixels -- proves the `src`-based
    # fast path in `_resolve_image_bytes` actually engaged, not just its
    # screenshot fallback (which would produce different, larger PNG bytes
    # for the same 1x1 source image once re-rendered/re-encoded).
    from app.core.blatt_kern_pptx_export_editable import extract_slide_elements

    width_emu, height_emu = _slide_size_emu("presentation_16_9")
    results = extract_slide_elements(rendered_image_html, width_emu, height_emu, mathjax_wait_ms=1000)

    assert len(results) == 1
    images = [el for el in (results[0].elements or []) if el.kind == "image"]
    assert any(img.image_bytes == _ONE_PX_PNG for img in images)


def test_extract_slide_elements_positions_stay_within_slide_bounds(rendered_presentation_html):
    from app.core.blatt_kern_pptx_export_editable import extract_slide_elements

    width_emu, height_emu = _slide_size_emu("presentation_16_9")
    results = extract_slide_elements(rendered_presentation_html, width_emu, height_emu, mathjax_wait_ms=3000)

    for result in results:
        for element in result.elements or []:
            assert 0 <= element.left_emu <= width_emu
            assert 0 <= element.top_emu <= height_emu
