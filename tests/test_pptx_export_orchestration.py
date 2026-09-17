"""Tests for `build_presentation_pptx`'s Stage A/B/C dispatch: does
`editable=True` actually try the editable path, does a per-slide
recoverable failure produce Stage B (rasterize just that slide, `PPTX001`)
without aborting the run, does a pipeline-level failure produce Stage C
(`PPTX002` + full raster fallback), and -- critically -- does an
UNEXPECTED bug propagate as a real error instead of silently becoming a
"successful" raster export.

All Playwright/browser I/O is mocked out (`extract_slide_elements`,
`build_worksheet_from_request`) -- this file tests the orchestration
logic in `blatt_kern_pptx_export.py`, not the extraction itself (covered
by `test_pptx_export_editable_convert.py` and the real-browser
integration test).
"""

from pathlib import Path

import pytest
from pptx import Presentation

from app.core.blatt_kern_pptx_export import build_presentation_pptx
from app.core.blatt_kern_pptx_export_editable import (
    EditableExportUnavailable,
    RenderableElement,
    SlideExtractionResult,
    TextRun,
)
from app.core.build_requests import WorksheetDesignOptions

_DESIGN = WorksheetDesignOptions("indigo", "segoe", "normal")

_ONE_PX_PNG = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
    b"\x08\x02\x00\x00\x00\x90wS\xde\x00\x00\x00\x0cIDATx\x9cc\xf8\xcf\xc0"
    b"\x00\x00\x03\x01\x01\x00\x18\xdd\x8d\xb0\x00\x00\x00\x00IEND\xaeB`\x82"
)


def _text_element(text="Hallo"):
    return RenderableElement(
        kind="text", left_emu=1000, top_emu=1000, width_emu=500_000, height_emu=100_000,
        align="left",
        runs=[TextRun(text=text, font_size_pt=12.0, bold=False, italic=False, color_rgb=(0, 0, 0))],
    )


def _stub_build_worksheet_from_request(monkeypatch):
    """Replaces the real HTML-rendering step with a no-op that just leaves
    the (already-created, empty) temp file in place -- the mocked
    `extract_slide_elements` never actually reads it."""

    monkeypatch.setattr(
        "app.core.blatt_kern_pptx_export.build_worksheet_from_request", lambda request: None
    )


def test_editable_true_uses_editable_path_and_writes_real_shapes(tmp_path, monkeypatch):
    _stub_build_worksheet_from_request(monkeypatch)
    monkeypatch.setattr(
        "app.core.blatt_kern_pptx_export.extract_slide_elements",
        lambda *_args, **_kwargs: [SlideExtractionResult(elements=[_text_element("Folie 1 Text")])],
    )

    out = tmp_path / "out.pptx"
    diagnostics = []
    build_presentation_pptx(
        Path("unused.md"), out, page_format="presentation_16_9", design=_DESIGN,
        editable=True, diagnostics_out=diagnostics,
    )

    assert out.exists()
    assert not any(d.code in {"PPTX001", "PPTX002"} for d in diagnostics)
    prs = Presentation(str(out))
    assert len(prs.slides) == 1
    shape = list(prs.slides[0].shapes)[0]
    assert shape.has_text_frame and shape.text_frame.text == "Folie 1 Text"


def test_stage_b_rasterizes_only_the_failed_slide_and_reports_pptx001(tmp_path, monkeypatch):
    _stub_build_worksheet_from_request(monkeypatch)
    monkeypatch.setattr(
        "app.core.blatt_kern_pptx_export.extract_slide_elements",
        lambda *_a, **_k: [
            SlideExtractionResult(elements=[_text_element("editierbare Folie")]),
            SlideExtractionResult(elements=None, raster_image_bytes=_ONE_PX_PNG, raster_reason="Timeout beim Screenshot"),
        ],
    )

    out = tmp_path / "out.pptx"
    diagnostics = []
    build_presentation_pptx(
        Path("unused.md"), out, page_format="presentation_16_9", design=_DESIGN,
        editable=True, diagnostics_out=diagnostics,
    )

    pptx001 = [d for d in diagnostics if d.code == "PPTX001"]
    assert len(pptx001) == 1
    assert "Folie 2" in pptx001[0].message
    assert "Timeout beim Screenshot" in pptx001[0].message

    prs = Presentation(str(out))
    assert len(prs.slides) == 2
    # Slide 1 stayed editable (real textbox); slide 2 is a single raster
    # picture -- both slides present, run did not abort.
    assert list(prs.slides[0].shapes)[0].has_text_frame
    assert not list(prs.slides[1].shapes)[0].has_text_frame


def test_stage_c_falls_back_to_full_raster_and_reports_pptx002(tmp_path, monkeypatch):
    _stub_build_worksheet_from_request(monkeypatch)
    monkeypatch.setattr(
        "app.core.blatt_kern_pptx_export.extract_slide_elements",
        lambda *_a, **_k: (_ for _ in ()).throw(EditableExportUnavailable("kein Browser gefunden")),
    )
    raster_calls = []
    monkeypatch.setattr(
        "app.core.blatt_kern_pptx_export._build_presentation_pptx_raster",
        lambda **kwargs: raster_calls.append(kwargs),
    )

    out = tmp_path / "out.pptx"
    diagnostics = []
    build_presentation_pptx(
        Path("unused.md"), out, page_format="presentation_16_9", design=_DESIGN,
        editable=True, diagnostics_out=diagnostics,
    )

    pptx002 = [d for d in diagnostics if d.code == "PPTX002"]
    assert len(pptx002) == 1
    assert "kein Browser gefunden" in pptx002[0].message
    assert len(raster_calls) == 1  # raster fallback actually ran


def test_editable_false_never_touches_the_editable_path(tmp_path, monkeypatch):
    called = []
    monkeypatch.setattr(
        "app.core.blatt_kern_pptx_export.extract_slide_elements",
        lambda *_a, **_k: called.append(True),
    )
    raster_calls = []
    monkeypatch.setattr(
        "app.core.blatt_kern_pptx_export._build_presentation_pptx_raster",
        lambda **kwargs: raster_calls.append(kwargs),
    )

    build_presentation_pptx(
        Path("unused.md"), tmp_path / "out.pptx", page_format="presentation_16_9",
        design=_DESIGN, editable=False,
    )

    assert called == []
    assert len(raster_calls) == 1


def test_unexpected_bug_in_editable_path_propagates_instead_of_silently_falling_back(tmp_path, monkeypatch):
    # A bug (here: a plain ValueError, standing in for e.g. a real coding
    # mistake) must NEVER be swallowed into an apparently-successful
    # raster export -- only EditableExportUnavailable is an accepted
    # Stage-C trigger.
    _stub_build_worksheet_from_request(monkeypatch)
    monkeypatch.setattr(
        "app.core.blatt_kern_pptx_export.extract_slide_elements",
        lambda *_a, **_k: (_ for _ in ()).throw(ValueError("bug: geometry math blew up")),
    )
    raster_calls = []
    monkeypatch.setattr(
        "app.core.blatt_kern_pptx_export._build_presentation_pptx_raster",
        lambda **kwargs: raster_calls.append(kwargs),
    )

    with pytest.raises(ValueError, match="bug: geometry math blew up"):
        build_presentation_pptx(
            Path("unused.md"), tmp_path / "out.pptx", page_format="presentation_16_9",
            design=_DESIGN, editable=True,
        )

    assert raster_calls == []  # no silent fallback for a real bug
